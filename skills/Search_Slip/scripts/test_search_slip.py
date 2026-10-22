import unittest
import os
import sys
import tempfile
import json
import subprocess
from openpyxl import load_workbook
import fitz
import numpy as np
from unittest.mock import patch

sys.path.insert(0, os.path.dirname(__file__))
from search_slip import clean_and_normalize_datetime, evidence_review_status, export_to_excel, extract_slip_details_ocr, search_slips_in_pdf


class TestSearchSlip(unittest.TestCase):
    def setUp(self):
        self.test_pdf = os.path.join(
            os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(__file__)))),
            "Folder_Out",
            "Evidence_Chat_V1.pdf"
        )

    def test_search_slips_on_v1(self):
        if not os.path.exists(self.test_pdf):
            self.skipTest(f"{self.test_pdf} does not exist")

        results = search_slips_in_pdf(self.test_pdf, output_excel=False, output_json=False)
        self.assertGreaterEqual(len(results), 2)
        
        # Check that page 64 is detected
        pages_detected = [p for item in results for p in item["pages"]]
        self.assertIn(64, pages_detected)
        self.assertIn(97, pages_detected)

    def test_decodes_reference_dates_for_current_year_not_only_2025(self):
        self.assertEqual(
            clean_and_normalize_datetime("", "2026060914451234"),
            "09 มิ.ย. 2569 - 14:45",
        )

    def test_ocr_extraction_is_labeled_for_human_review(self):
        self.assertEqual(evidence_review_status(), "OCR candidate - ต้องตรวจทาน")

    def test_excel_export_marks_ocr_fields_as_review_candidates(self):
        item = {
            "pages": [1], "datetime": "-", "sender_bank": "-", "sender_name": "-",
            "amount": "5000.00", "receiver_name": "-", "receiver_bank": "-",
            "memo": "-", "ref_id": "ref-test",
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = os.path.join(temp_dir, "slips_Slip_Index.xlsx")
            export_to_excel([item], output_path)
            workbook = load_workbook(output_path, read_only=True)
            self.assertEqual(workbook.active["J3"].value, "OCR candidate - ต้องตรวจทาน")
            workbook.close()

    def test_retains_visual_slip_candidate_when_ocr_returns_no_fields(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = os.path.join(temp_dir, "one_page.pdf")
            document = fitz.open()
            document.new_page(width=300, height=400)
            document.save(pdf_path)
            document.close()
            with patch("search_slip.detect_slip_in_image", return_value=(True, (20, 20, 100, 160))), patch(
                "search_slip.extract_slip_details_ocr", return_value={}
            ):
                results = search_slips_in_pdf(pdf_path, output_excel=True, output_json=True)
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["pages"], [1])
        self.assertEqual(results[0]["review_status"], "OCR candidate - ต้องตรวจทาน")

    def test_scans_password_protected_pdf_with_supplied_password(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = os.path.join(temp_dir, "encrypted.pdf")
            document = fitz.open()
            document.new_page(width=300, height=400)
            document.save(pdf_path, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="case-pass")
            document.close()
            results = search_slips_in_pdf(pdf_path, output_excel=False, output_json=False, password="case-pass")
        self.assertEqual(results, [])

    def test_api_cli_writes_review_results_only_to_requested_temp_path(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = os.path.join(temp_dir, "evidence.pdf")
            output_path = os.path.join(temp_dir, "result.json")
            document = fitz.open()
            document.new_page(width=300, height=400)
            document.save(pdf_path)
            document.close()
            request = json.dumps({"pdf_path": pdf_path, "output_path": output_path, "password": ""})
            completed = subprocess.run(
                [sys.executable, os.path.join(os.path.dirname(__file__), "search_slip.py"), "--api-json-stdin"],
                input=request,
                capture_output=True,
                text=True,
                check=True,
            )
            self.assertEqual(json.loads(completed.stdout)["candidate_count"], 0)
            with open(output_path, encoding="utf-8") as result_file:
                self.assertEqual(json.load(result_file)["results"], [])
            self.assertFalse(os.path.exists(os.path.join(temp_dir, "evidence_Slip_Index.json")))
            self.assertFalse(os.path.exists(os.path.join(temp_dir, "evidence_Slip_Index.xlsx")))

    def test_ocr_uses_thai_script_model_when_tha_language_is_unavailable(self):
        with patch("search_slip.pytesseract.image_to_string", side_effect=[RuntimeError("missing tha"), "ภาษาไทย"]) as ocr:
            result = extract_slip_details_ocr(np.zeros((32, 32, 3), dtype=np.uint8))
        self.assertEqual(result["raw_text"], "ภาษาไทย")
        self.assertEqual([call.kwargs["lang"] for call in ocr.call_args_list], ["tha+eng", "script/Thai+eng"])

    def test_empty_ocr_retries_once_with_enlarged_sparse_text_crop(self):
        with patch("search_slip.pytesseract.image_to_string", side_effect=["", "", "", "7,200.00"]) as ocr:
            result = extract_slip_details_ocr(np.zeros((32, 32, 3), dtype=np.uint8))

        self.assertEqual(result["amount"], "7,200.00")
        self.assertEqual(len(ocr.call_args_list), 4)
        self.assertEqual(ocr.call_args_list[-1].args[0].shape, (64, 64))
        self.assertEqual(ocr.call_args_list[-1].kwargs["config"], "--psm 11")

    def test_extracts_date_when_thai_script_ocr_spaces_month_glyphs(self):
        raw_text = "วันท ี ่ ท ํ า ร า ย ก า ร 09 ม ิ . ย . 2568 - 14:45"
        with patch("search_slip.pytesseract.image_to_string", return_value=raw_text):
            result = extract_slip_details_ocr(np.zeros((32, 32, 3), dtype=np.uint8))
        self.assertEqual(result["datetime"], "09 มิ.ย. 2568 - 14:45")

    def test_ref_timestamp_does_not_override_conflicting_visible_slip_time(self):
        raw_text = """รหัสอ้างอิง 202506171502627774
17 บ ิ ย . 68, 15:51 น .
จำนวนเงิน 7,200.00 บาท"""
        with patch("search_slip.pytesseract.image_to_string", return_value=raw_text):
            result = extract_slip_details_ocr(np.zeros((32, 32, 3), dtype=np.uint8))

        self.assertEqual(result["datetime"], "-")
        self.assertIn("15:51", result["raw_text"])

    def test_assigns_spaced_ocr_bank_names_to_sender_and_receiver_separately(self):
        raw_text = """รหัสอ้างอิง Afelcd222498c42a5
ส ิ บ ต ร ี ณั ฐ ช ั ย ร
ก ร ุ ง ไท ย
XXX-X-XX526-6
ไป ย ั ง
น . ส . จ ิ ณ ห ์ น ิ ภ า ป ร ะ ส า ท เข ต ร ก า ร
ท ี เอ ็ ม บ ี ธน ชา ต
XXX-X-XX996-5
จ ํ า น ว น เง ิ น 5,000.00 บ า ท
"""
        with patch("search_slip.pytesseract.image_to_string", return_value=raw_text):
            result = extract_slip_details_ocr(np.zeros((32, 32, 3), dtype=np.uint8))

        self.assertEqual(result["sender_bank"], "กรุงไทย")
        self.assertEqual(result["receiver_bank"], "ทีเอ็มบีธนชาต (ttb)")
        self.assertEqual(result["raw_text"], raw_text)

    def test_recovers_party_candidates_from_account_blocks_when_transfer_labels_are_missing(self):
        raw_text = """ร ห ั ส อ ้ า ง อ ิ ง Afelcd222498c42a5
ส ิ บ ต ร ี ณั ฐ ช ั ย ร
ก ร ุ ง ไท ย
XXX-X-XX526-6
ub ESS
น . ส . จ ิ ณ ห ์ น ิ ภา ป ร ะ ส า ท เข ต ร ท กา ร
ที เอ ็ ม บ ี ธน ชา ต
XXX-X-XX996-5
จ ํ า น ว น เง ิ น 5,000.00 vn"""
        with patch("search_slip.pytesseract.image_to_string", return_value=raw_text):
            result = extract_slip_details_ocr(np.zeros((32, 32, 3), dtype=np.uint8))

        self.assertIn("ณัฐชัย", result["sender"])
        self.assertIn("XXX-X-XX526-6", result["sender"])
        self.assertIn("จิณห์นิภา", result["receiver"])
        self.assertIn("XXX-X-XX996-5", result["receiver"])
        self.assertEqual(result["sender_bank"], "กรุงไทย")
        self.assertEqual(result["receiver_bank"], "ทีเอ็มบีธนชาต (ttb)")

    def test_noisy_from_icon_is_not_a_sender_and_banks_follow_account_blocks(self):
        raw_text = """© จ า ก
ส ิ บ ต ร ี ณั ฐ ช ั ย s***
ก ร ุ ง ไท ย
XXX-X-XX526-6
ไป ย ฮั ง
น . ส . จ ิ ณ ห ์ น ิ ภา ป ร ะ ส า ท เข ต ร ก า ร
ท ี เอ ็ บ ม ี ธน ชา ต
XXX-X-XX996-5
จ ํ า น ว น เง ิ น 3,000.00 vn"""
        with patch("search_slip.pytesseract.image_to_string", return_value=raw_text):
            result = extract_slip_details_ocr(np.zeros((32, 32, 3), dtype=np.uint8))

        self.assertIn("ณัฐชัย", result["sender"])
        self.assertNotIn("©", result["sender"])
        self.assertEqual(result["sender_bank"], "กรุงไทย")
        self.assertEqual(result["receiver_bank"], "ทีเอ็มบีธนชาต (ttb)")

    def test_unlinked_chat_name_does_not_fill_account_block(self):
        raw_text = """นาย คนในข้อความสนทนา
XXX-X-XX526-6
กรุงไทย
XXX-X-XX996-5
จำนวนเงิน 5,000.00 บาท"""
        with patch("search_slip.pytesseract.image_to_string", return_value=raw_text):
            result = extract_slip_details_ocr(np.zeros((32, 32, 3), dtype=np.uint8))

        self.assertEqual(result["sender"], "ไม่ระบุชื่อ (XXX-X-XX526-6)")
        self.assertEqual(result["receiver"], "ไม่ระบุชื่อ (XXX-X-XX996-5)")

    def test_masked_account_suffix_does_not_invent_party_names(self):
        raw_text = "จากบัญชี\nกรุงไทย\nXXX-X-XX526-6\nไปยัง\nทีทีบี\nXXX-X-XX996-5\nจำนวนเงิน 5,000.00 บาท"
        with patch("search_slip.pytesseract.image_to_string", return_value=raw_text):
            result = extract_slip_details_ocr(np.zeros((32, 32, 3), dtype=np.uint8))

        self.assertEqual(result["sender"], "ไม่ระบุชื่อ (XXX-X-XX526-6)")
        self.assertEqual(result["receiver"], "ไม่ระบุชื่อ (XXX-X-XX996-5)")

    def test_extracts_party_names_next_to_transfer_labels(self):
        raw_text = """จาก
ธนาคารกรุงไทย
นาย สมชาย ใจดี
XXX-X-XX123-4
ไปยัง
น.ส. สมหญิง รักดี
ธนาคารไทยพาณิชย์
XXX-X-XX567-8
จำนวนเงิน 5,000.00 บาท"""
        with patch("search_slip.pytesseract.image_to_string", return_value=raw_text):
            result = extract_slip_details_ocr(np.zeros((32, 32, 3), dtype=np.uint8))

        self.assertIn("สมชาย ใจดี", result["sender"])
        self.assertIn("สมหญิง รักดี", result["receiver"])
        self.assertIn("XXX-X-XX123-4", result["sender"])
        self.assertIn("XXX-X-XX567-8", result["receiver"])
        self.assertEqual(result["raw_text"], raw_text)

    def test_bank_label_is_not_promoted_to_party_name(self):
        raw_text = "จากบัญชี\nธนาคารกรุงไทย\nXXX-X-XX526-6\nไปยัง\nธนาคารไทยพาณิชย์\nXXX-X-XX996-5\nจำนวนเงิน 5,000.00 บาท"
        with patch("search_slip.pytesseract.image_to_string", return_value=raw_text):
            result = extract_slip_details_ocr(np.zeros((32, 32, 3), dtype=np.uint8))

        self.assertEqual(result["sender"], "ไม่ระบุชื่อ (XXX-X-XX526-6)")
        self.assertEqual(result["receiver"], "ไม่ระบุชื่อ (XXX-X-XX996-5)")

    def test_does_not_promote_unlinked_person_name_to_sender(self):
        raw_text = """นาย คนในข้อความสนทนา
จากบัญชี
ธนาคารกรุงไทย
XXX-X-XX526-6
ไปยัง
ธนาคารไทยพาณิชย์
XXX-X-XX996-5
จำนวนเงิน 5,000.00 บาท"""
        with patch("search_slip.pytesseract.image_to_string", return_value=raw_text):
            result = extract_slip_details_ocr(np.zeros((32, 32, 3), dtype=np.uint8))

        self.assertEqual(result["sender"], "ไม่ระบุชื่อ (XXX-X-XX526-6)")
        self.assertEqual(result["receiver"], "ไม่ระบุชื่อ (XXX-X-XX996-5)")

    def test_extracts_party_names_from_spaced_ocr_and_noisy_labels(self):
        raw_text = """อา ก
ส ิ บ ต ร ี ณั ฐ ช ั ย s***
ก ร ุ ง ไท ย
XXX-X-XX526-6
ไป ฮั ง
น . ส . จ ิ ณ ห ์ นิ ภา ป ร ะ ส า ท เข ต ก า ร
ท ี เอ ็ ม บ ี ธน ชา ต
XXX-X-XX996-5
จ ํ า น ว น เข ิ น 3,000.00 vn"""
        with patch("search_slip.pytesseract.image_to_string", return_value=raw_text):
            result = extract_slip_details_ocr(np.zeros((32, 32, 3), dtype=np.uint8))

        self.assertIn("ณัฐชัย", result["sender"])
        self.assertIn("จิณห์นิภา", result["receiver"].replace(" ", ""))
        self.assertIn("XXX-X-XX526-6", result["sender"])
        self.assertIn("XXX-X-XX996-5", result["receiver"])
        self.assertEqual(result["raw_text"], raw_text)

    def test_preserves_each_page_occurrence_when_transaction_reference_repeats(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = os.path.join(temp_dir, "repeated_ref.pdf")
            document = fitz.open()
            document.new_page(width=300, height=400)
            document.new_page(width=300, height=400)
            document.save(pdf_path)
            document.close()
            details_by_page = iter([
                {
                "ref_id": "same-reference-123456", "amount": "5000.00",
                "datetime": "09 มิ.ย. 2568 - 14:45", "sender": "ไม่ระบุชื่อ (อ่านจากภาพไม่ได้)",
                "receiver": "ไม่ระบุชื่อ (อ่านจากภาพไม่ได้)", "sender_bank": "-",
                "receiver_bank": "-", "memo": "-", "raw_text": "OCR page one",
                },
                {
                    "ref_id": "same-reference-123456", "amount": "5000.00",
                    "datetime": "09 มิ.ย. 2568 - 14:45", "sender": "ไม่ระบุชื่อ (อ่านจากภาพไม่ได้)",
                    "receiver": "ไม่ระบุชื่อ (อ่านจากภาพไม่ได้)", "sender_bank": "-",
                    "receiver_bank": "-", "memo": "-", "raw_text": "OCR page two",
                },
            ])
            with patch("search_slip.detect_slip_in_image", return_value=(True, (20, 20, 100, 160))), patch(
                "search_slip.extract_slip_details_ocr", side_effect=lambda _crop: next(details_by_page)
            ):
                results = search_slips_in_pdf(pdf_path, output_excel=True, output_json=True)

            self.assertEqual(len(results), 2)
            self.assertEqual([item["pages"] for item in results], [[1], [2]])
            self.assertEqual([item["occurrence_id"] for item in results], ["pdf-page-1", "pdf-page-2"])
            self.assertEqual([item["raw_text"] for item in results], ["OCR page one", "OCR page two"])
            self.assertFalse(results[0]["repeated_ref"])
            self.assertTrue(results[1]["repeated_ref"])
            self.assertIn("ต้องตรวจว่าเป็นรายการซ้ำ", results[1]["review_status"])
            with open(os.path.join(temp_dir, "repeated_ref_Slip_Index.json"), encoding="utf-8") as result_file:
                exported = json.load(result_file)
            self.assertEqual([entry["occurrence_id"] for entry in exported], ["pdf-page-1", "pdf-page-2"])
            self.assertEqual([entry["repeated_ref"] for entry in exported], [False, True])
            workbook = load_workbook(os.path.join(temp_dir, "repeated_ref_Slip_Index.xlsx"), read_only=True)
            self.assertIn("ต้องตรวจว่าเป็นรายการซ้ำ", workbook.active["J4"].value)
            workbook.close()

    def test_preserves_separate_page_occurrences_when_reference_is_missing(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = os.path.join(temp_dir, "missing_ref.pdf")
            document = fitz.open()
            document.new_page(width=300, height=400)
            document.new_page(width=300, height=400)
            document.save(pdf_path)
            document.close()
            details_by_page = iter([
                {"raw_text": "OCR page one", "ref_id": "", "amount": "5000.00"},
                {"raw_text": "OCR page two", "ref_id": "", "amount": "5000.00"},
            ])
            with patch("search_slip.detect_slip_in_image", return_value=(True, (20, 20, 100, 160))), patch(
                "search_slip.extract_slip_details_ocr", side_effect=lambda _crop: next(details_by_page)
            ):
                results = search_slips_in_pdf(pdf_path, output_excel=False, output_json=False)

        self.assertEqual([item["occurrence_id"] for item in results], ["pdf-page-1", "pdf-page-2"])
        self.assertEqual([item["pages"] for item in results], [[1], [2]])
        self.assertEqual([item["raw_text"] for item in results], ["OCR page one", "OCR page two"])
        self.assertFalse(any(item["repeated_ref"] for item in results))



if __name__ == "__main__":
    unittest.main()
